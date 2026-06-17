"""Gera eval/pool_revisao.xlsx a partir do TSV, pronto para anotacao:
- coluna rel_revisado com lista suspensa (0/1/2), pre-preenchida com a sugestao
- destaque nas linhas criticas (sim_densa alta + overlap baixo)
- cabecalho fixo, filtro, larguras e quebra de texto
- aba de instrucoes
"""
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

SRC = "eval/pool_revisao.tsv"
OUT = "eval/pool_revisao.xlsx"

with open(SRC, encoding="utf-8") as f:
    rows = list(csv.reader(f, delimiter="\t"))
header, data = rows[0], rows[1:]

wb = Workbook()
ws = wb.active
ws.title = "anotacao"

head_fill = PatternFill("solid", fgColor="305496")
head_font = Font(bold=True, color="FFFFFF")
crit_fill = PatternFill("solid", fgColor="FFF2CC")  # amarelo claro
rev_fill = PatternFill("solid", fgColor="E2EFDA")   # verde claro (coluna a preencher)
thin = Side(style="thin", color="D9D9D9")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

# cabecalho
for j, col in enumerate(header, 1):
    c = ws.cell(1, j, col)
    c.fill = head_fill; c.font = head_font
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = border

col = {name: i + 1 for i, name in enumerate(header)}
crev = col["rel_revisado"]; csug = col["rel_sugerido"]

# dados (pre-preenche rel_revisado com a sugestao)
for i, r in enumerate(data, start=2):
    for j, val in enumerate(r, 1):
        cell = ws.cell(i, j, val)
        cell.border = border
        cell.alignment = Alignment(vertical="top", wrap_text=(header[j-1] in ("query", "titulo", "abstract")))
    rv = ws.cell(i, crev, r[csug - 1] if r[csug - 1] != "" else "")
    rv.fill = rev_fill
    rv.alignment = Alignment(horizontal="center")

# numericos de fato numeros
for name in ("sim_densa", "overlap", "rel_sugerido", "rel_revisado"):
    cc = col[name]
    for i in range(2, len(data) + 2):
        try:
            ws.cell(i, cc).value = float(ws.cell(i, cc).value)
        except (TypeError, ValueError):
            pass

# lista suspensa 0/1/2 na coluna rel_revisado
dv = DataValidation(type="list", formula1='"0,1,2"', allow_blank=True, showErrorMessage=True)
dv.error = "Use 0, 1 ou 2"; dv.prompt = "0=nao relevante, 1=relevante, 2=muito relevante"
ws.add_data_validation(dv)
L = get_column_letter(crev)
dv.add(f"{L}2:{L}{len(data)+1}")

# destaque das linhas criticas (sim_densa>=0.5 e overlap<=0.3) na coluna rel_revisado
fsim = get_column_letter(col["sim_densa"]); fov = get_column_letter(col["overlap"])
ws.conditional_formatting.add(
    f"{L}2:{L}{len(data)+1}",
    FormulaRule(formula=[f"AND(${fsim}2>=0.5,${fov}2<=0.3)"], fill=crit_fill))

# larguras
widths = {"qid": 6, "query": 38, "doc_id": 22, "sistemas": 14, "fonte": 10,
          "sim_densa": 10, "overlap": 9, "rel_sugerido": 11, "rel_revisado": 12,
          "titulo": 45, "abstract": 70}
for name, w in widths.items():
    ws.column_dimensions[get_column_letter(col[name])].width = w

ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{len(data)+1}"

# aba instrucoes
ins = wb.create_sheet("instrucoes")
linhas = [
    ("Como anotar", ""),
    ("Preencha a coluna rel_revisado (aba anotacao) com 0, 1 ou 2.", ""),
    ("", ""),
    ("2 = muito relevante", "trata do cruzamento tecnica + dominio da consulta"),
    ("1 = relevante", "trata da tecnica OU do dominio, nao dos dois"),
    ("0 = nao relevante", "nao trata de nenhum, ou so tangencia"),
    ("", ""),
    ("Linhas em AMARELO", "sim_densa alta e overlap baixo: provavel relevante mesmo se a sugestao for 0 (caso cross-lingual). Olhe com atencao."),
    ("rel_sugerido", "palpite automatico (acerta ~52%). Ja veio copiado em rel_revisado; ajuste o que estiver errado."),
    ("sim_densa", "proximidade de significado (0 a 1)."),
    ("overlap", "fracao de palavras da query no documento (0 a 1)."),
    ("", ""),
    ("Dica", "ordenado por consulta e por sim_densa. Quando a sim cai muito, costuma ser tudo 0."),
]
for i, (a, b) in enumerate(linhas, 1):
    ca = ins.cell(i, 1, a); ca.font = Font(bold=(b == "" and a != ""))
    ins.cell(i, 2, b)
ins.column_dimensions["A"].width = 26
ins.column_dimensions["B"].width = 90

wb.save(OUT)
print("gerado:", OUT, "| linhas:", len(data))
